#!/usr/bin/env python3
"""
==========================================================
MEDIVOX — scientific presentation deck (English, 16:9)
----------------------------------------------------------
Built from the same slide renderer as the scientific video
(tools/bangun-video-ilmiah.py), so the deck, the video and the
web app share one neumorphic design language. Slides are
rendered at their final, fully-animated state; Blender stills
come from the showreel shots. Every slide carries full speaker
notes, so the talk can be given straight from the deck.

Run:     python tools/bangun-deck.py
Output:  build/MEDIVOX_Deck.pptx
==========================================================
"""
import os, sys, glob, importlib.util
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

AKAR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
GBR = os.path.join(AKAR, 'build', 'deck', 'gbr')
KELUAR = os.path.join(AKAR, 'build', 'MEDIVOX_Deck.pptx')
os.makedirs(GBR, exist_ok=True)

_spec = importlib.util.spec_from_file_location('vi', os.path.join(AKAR, 'tools', 'bangun-video-ilmiah.py'))
VI = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(VI)


def akhir(fn, n=120, pada=0.93):
    """Render a slide function at the given point of its timeline."""
    return fn(int(n * pada), n)


def dari_gen(gen_fn, pada=0.75, n=160):
    g = gen_fn()
    im = None
    for k in range(int(n * pada)):
        im = next(g)
    return im


def shot(nama, eyebrow, judul, sub, pada=0.6, overlay=None, bab=''):
    a = VI.shot(nama, 5.0, eyebrow, judul, sub, bab, overlay)
    return a.fn(int(a.n * pada), a.n)


def dengan_kepala(im, bab):
    VI.BAB_AKTIF[0] = bab
    r = im.convert('RGBA')
    VI.kepala(r)
    return r.convert('RGB')


def bab(no, label, judul, sub):
    return VI.s_bab(no, label, judul, sub).fn(96, 102)


# (image builder, speaker notes)
SLIDES = [
    (lambda: akhir(VI.s_judul),
     'MEDIVOX is a research prototype: a pseudo-holographic, touchless display for DICOM volumes. '
     'It renders CT and MRI data in the browser, projects four views into a glass prism, and is controlled by voice '
     'and hand gestures. Nothing is installed and no image data is uploaded.'),
    (lambda: dengan_kepala(akhir(VI.s_agenda), 'Overview'),
     'Eight parts: why this problem matters, what already exists, the gap we target and our contribution, how the system '
     'is built, its features including voice and gesture control, how it is used, how it is tested, and its limits.'),
    (lambda: bab('01', 'Background', 'Anatomy is 3D. Screens are flat.', 'Why reading volumes slice by slice is still hard, and hard to share.'),
     'Section one: background.'),
    (lambda: dengan_kepala(akhir(VI.s_irisan), 'Background'),
     'A CT or MRI study is a stack of hundreds of flat images. Readers scroll through them and rebuild the anatomy '
     'in their head. That skill takes years to build and is hard to hand over to anyone else in the room.'),
    (lambda: dengan_kepala(akhir(VI.s_angka), 'Background'),
     'A typical CT volume is around 512 by 512 by several hundred voxels; a study means scrolling hundreds of slices on a '
     'flat screen, usually by one person at a time. These numbers describe the data, not clinical performance.'),
    (lambda: dengan_kepala(akhir(VI.s_masalah), 'Background'),
     'Two practical gaps. First, shared understanding: in a tumour board only the reader truly sees the 3D anatomy. '
     'Second, sterile or hands-busy settings, where touching a mouse or keyboard is a contamination risk.'),
    (lambda: dengan_kepala(akhir(VI.s_pertanyaan), 'Background'),
     'Our research question: can a low-cost, glasses-free volumetric display, controlled without touch and fed directly '
     'from DICOM in a browser, make 3D imaging easier to see, share and discuss?'),
    (lambda: bab('02', 'State of the art', 'What exists today', 'From flat PACS viewers to headsets, printed models and light-field panels.'),
     'Section two: the state of the art.'),
    (lambda: dengan_kepala(akhir(VI.s_matriks, pada=0.98), 'State of the art'),
     'A qualitative comparison, made by us, not a benchmark. Flat viewers are the clinical standard but give no real depth. '
     'Headsets give depth but isolate each viewer. Printed models are tangible but slow and static. Light-field panels are '
     'glasses-free but costly. Consumer Pepper’s-ghost pyramids are cheap but play videos, not DICOM. MEDIVOX targets the '
     'combination, with an honest caveat: it produces a pseudo-3D floating image, not a true volumetric light field.'),
    (lambda: dengan_kepala(akhir(VI.s_kurang), 'State of the art'),
     'Where each option falls short: flat viewers ask the reader to imagine depth, headsets need one device per person and '
     'cut people off from the room, and printed models take hours and show one frozen case.'),
    (lambda: bab('03', 'Gap & novelty', 'The missing middle', 'A shared, touchless 3D view that works from raw DICOM on ordinary hardware.'),
     'Section three: the research gap and our contribution.'),
    (lambda: dengan_kepala(akhir(VI.s_venn, pada=0.97), 'Gap & novelty'),
     'The gap is the intersection of three properties that existing systems do not combine: a glasses-free view that '
     'several people can share, touchless control, and working straight from DICOM in a browser.'),
    (lambda: dengan_kepala(akhir(VI.s_novelty), 'Gap & novelty'),
     'Four contributions: a browser-native pipeline with our own DICOM parser and CPU ray casting; four-view rendering '
     'tuned for any prism; multimodal touchless control where voice and gestures share one command path; and commodity '
     'hardware, from a tablet with an acrylic pyramid to the MEDIVOX-1 device.'),
    (lambda: bab('04', 'System design', 'How MEDIVOX works', 'Data path, optics, hardware and rendering.'),
     'Section four: system design.'),
    (lambda: dengan_kepala(akhir(VI.s_arsitektur, pada=0.97), 'System design'),
     'Architecture. DICOM files are read with the File API, parsed by our own parser, turned into a voxel volume at true '
     'millimetre spacing, and ray cast. Four views are laid out around the centre, shown on the emitter, and reflected by '
     'the prism into a floating image. Voice and gesture enter through a single controller. Everything runs in the browser '
     'on the CPU; only speech recognition may use a cloud service.'),
    (lambda: dengan_kepala(akhir(VI.s_optik), 'System design'),
     'Optics. Each glass face at 45 degrees reflects one view; because the glass is nearly clear, the reflection appears to '
     'hang behind it. Four faces and four views 90 degrees apart make the form hold together from every side. There is no '
     'eye tracking and no glasses; depth is suggested rather than stereoscopic.'),
    (lambda: shot('ledak', 'Hardware', 'MEDIVOX-1, exploded', 'Prism, emitter, microphone array and camera window',
                  pada=0.92, overlay=VI.overlay_ledak),
     'The MEDIVOX-1 hardware, exploded: the floating volume, the glass prism, the four-view emitter panel, a sixteen-port '
     'microphone array with a listening light ring, and a camera window for hand tracking above the compute base.'),
    (lambda: dengan_kepala(akhir(VI.s_render, pada=0.98), 'System design'),
     'Rendering. Rays are cast through the volume from 24 angles once, then the frames are replayed, so rotation stays '
     'smooth even on a CPU. The right side shows the real emitter output of the app: four views around a centre mark.'),
    (lambda: bab('05', 'Features', 'Features', 'The device, the apps, and touchless control.'),
     'Section five: features.'),
    (lambda: shot('orbit', 'Pepper’s ghost', 'Four reflections, one floating form.', 'No glasses, no projector, no headset'),
     'The device in the studio: one white base, one glass prism, and a volume that appears to float inside it.'),
    (lambda: dengan_kepala(akhir(VI.s_grid), 'Features'),
     'Eight capabilities: the hologram stage, a complete 2D viewer, MPR, MIP and 3D series, surface export to STL and OBJ, '
     'voice commands, hand gestures, an anatomy atlas, and privacy by design.'),
    (lambda: shot('ekosistem', 'Ecosystem', 'One study, every screen.', 'MEDIVOX-1, laptop and Android phone read the same study'),
     'One study on every screen: the MEDIVOX-1 device, a laptop and an Android phone all run the same web app.'),
    (lambda: dari_gen(VI.SR.adegan_alur, pada=0.92),
     'The workflow: pick a study in the worklist, read it in the 2D viewer, and send it to the hologram stage. DICOM is parsed '
     'in the browser and nothing is uploaded.'),
    (lambda: dari_gen(VI.SR.adegan_android, pada=0.8),
     'Mobile-first: on Android the app uses a bottom navigation with a central hologram button, drawers for panels, and '
     'controls within thumb reach.'),
    (lambda: shot('suara', 'Voice control', 'Say it, and it happens.', 'Spoken command › grammar › the same controls as the buttons',
                  pada=0.36, overlay=VI.overlay_suara),
     'Voice control in use. The clinician says “rotate left”; the command is recognised and the hologram turns. '
     'The microphone ring on the device lights up while it listens.'),
    (lambda: dengan_kepala(VI.s_pipa('Voice pipeline', [
        ('mik', 'Microphone', 'Speech captured only after the user turns voice on'),
        ('awan', 'Web Speech API', 'Browser recogniser returns a transcript (cloud in Chrome)'),
        ('teks', 'Command grammar', 'Longest-match lookup over a fixed EN + ID vocabulary'),
        ('setelan', 'Controller', 'Same handlers as buttons, chips and keys'),
        ('hologram', 'Hologram', 'Rotate, zoom, change mode, pick an organ')],
        'A small, closed vocabulary keeps recognition reliable and makes every command predictable.', 'Voice control')(110, 120), 'Features'),
     'The voice pipeline. The browser’s speech recogniser returns a transcript; in Chrome this uses Google’s cloud service, '
     'which the app states before the microphone starts. A longest-match grammar over a fixed English and Indonesian '
     'vocabulary maps it to a command, handled by the same code as the on-screen buttons.'),
    (lambda: shot('gestur', 'Gesture control', 'Swipe to turn it.', 'The camera follows the hand; the hologram follows the camera',
                  pada=0.42, overlay=VI.overlay_gestur),
     'Gesture control in use. The hand is tracked by the camera; moving it left or right turns the hologram, and moving it '
     'closer or further changes its size.'),
    (lambda: dengan_kepala(VI.s_pipa('Gesture pipeline', [
        ('kamera', 'Camera frame', '15 fps, low resolution, never stored'),
        ('tangan', 'Skin & motion mask', 'Pixels that look like a moving hand'),
        ('probe', 'Centroid & box', 'Hand position and apparent size'),
        ('sinkron', 'Smoothing', 'Exponential filter removes jitter'),
        ('hologram', 'Rotation & scale', 'Left-right turns it, near-far resizes it')],
        'Frames are processed in the browser and discarded immediately. Nothing leaves the device.', 'Gesture control')(110, 120), 'Features'),
     'The gesture pipeline: low-resolution camera frames at 15 frames per second, a skin-and-motion mask, the hand’s '
     'centroid and box, an exponential smoothing filter, then rotation and scale. Frames are discarded after processing.'),
    (lambda: bab('06', 'In use', 'In the room', 'A floating image the whole team can gather around.'),
     'Section six: MEDIVOX in use.'),
    (lambda: shot('tim', 'Tumour board', 'Everyone sees the same thing.', 'Each person looks from their own side of the table', pada=0.7),
     'A tumour board around MEDIVOX-1: each person looks at the same anatomy from their own side of the table, while one '
     'of them points and another talks it through.'),
    (lambda: dengan_kepala(akhir(VI.s_guna), 'In use'),
     'Use cases: tumour boards, surgical planning, sterile settings, teaching, patient conversations, and privacy-sensitive '
     'environments where data must stay on the device.'),
    (lambda: bab('07', 'Validation', 'How it is checked', 'Automated tests, real DICOM data, and what is still missing.'),
     'Section seven: validation.'),
    (lambda: dengan_kepala(akhir(VI.s_validasi), 'Validation'),
     'What is tested: 127 unit tests for the parser, volume, surface, models and voice grammar; 79 smoke tests that run the '
     'real pages in a DOM double; and real DICOM files from pydicom-data plus a TCIA series. No pixel data is uploaded. '
     'What is not done yet: a user study with clinicians and any clinical accuracy evaluation.'),
    (lambda: bab('08', 'Limitations & future work', 'What it is not, yet', 'The honest limits of a research prototype, and where it goes next.'),
     'Section eight: limitations and future work.'),
    (lambda: dengan_kepala(akhir(VI.s_batas), 'Limitations & future work'),
     'Limits: the image is pseudo-3D, it depends on room light and angle, voice uses a cloud service in Chrome, gesture '
     'tracking is colour-based, and there is no clinical validation. Next: on-device speech, learned hand landmarks, '
     'JPEG 2000 and DICOMweb, automatic prism calibration, and a structured user study.'),
    (lambda: akhir(VI.s_penutup, pada=0.8),
     'In short: a shared, glasses-free 3D view of DICOM volumes, controlled by voice and gesture, rendered in the browser '
     'with pixels that never leave the device. Try it at medivox-id.web.app. Research prototype, not a medical device.'),
]


def bangun():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    kosong = prs.slide_layouts[6]
    for k, (buat, catatan) in enumerate(SLIDES):
        im = buat()
        jalur = os.path.join(GBR, 's%02d.jpg' % (k + 1))
        im.save(jalur, 'JPEG', quality=92, subsampling=0)
        s = prs.slides.add_slide(kosong)
        s.shapes.add_picture(jalur, 0, 0, prs.slide_width, prs.slide_height)
        s.notes_slide.notes_text_frame.text = catatan
        print('  slide %2d' % (k + 1), flush=True)
    prs.save(KELUAR)
    print('deck: %s (%d slides, %.1f MB)' % (KELUAR, len(SLIDES), os.path.getsize(KELUAR) / 1048576))


if __name__ == '__main__':
    bangun()
